import io
from pptx import Presentation
from pptx.util import Inches, Pt
import re
import streamlit as st
from result import text_to_speech, translate_text
import time


LANGUAGE_CODES = {
    "Hindi": "hi-IN",
    "Bengali": "bn-IN",
    "Kannada": "kn-IN",
    "Malayalam": "ml-IN",
    "Marathi": "mr-IN",
    "Odia": "od-IN",
    "Punjabi": "pa-IN",
    "Tamil": "ta-IN",
    "Telugu": "te-IN",
    "Gujarati": "gu-IN"
}



background_image_path = "background.jpg"

def show_ppt():
    st.title("PowerPoint Outline")
    
    
    ppt_text_body = st.session_state.ppt_text_body
    
    
    if 'transcript_text_body' not in st.session_state:
        st.error("No transcript content found. Please go back to the result page and generate the content.")
        return

    transcript_text_body = st.session_state.transcript_text_body
    slides = re.split(r'Slide \d+:', transcript_text_body)[1:]

    if st.session_state.tts_option == "Yes" and st.session_state.language != "English":
        target_language = LANGUAGE_CODES[st.session_state.language]        
        for i, slide_content in enumerate(slides, 1):
            st.subheader(f"Slide {i}")
            st.write(slide_content.strip())

            sentences = re.split(r'(?<=[.!?])\s+', slide_content.strip())  # Split slide content into sentences
            
            st.subheader(f"Translated Slide {i} ({target_language})")
            
            for sentence in sentences:
                with st.spinner(f"Translating and generating audio for sentence: {sentence}"):
                    translated_text = translate_text(sentence, target_language)
                    st.write(translated_text)
                    
                    audio_bytes = text_to_speech(translated_text, target_language)  
                    
                    if audio_bytes:
                        st.audio(audio_bytes, format="audio/mp3")
                    else:
                        st.error(f"Failed to generate audio for the sentence: {sentence}")
                    
                    time.sleep(2)
    else:
        for i, slide_content in enumerate(slides, 1):
            st.subheader(f"Slide {i}")
            st.write(slide_content.strip())
    ppt_bytes = create_presentation(ppt_text_body, "background.jpg")

    st.download_button(
                label="Download Presentation",
                data=ppt_bytes,
                file_name="Python_Introduction.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
    
    if st.button("Back to Result"):
        st.session_state.current_page = "result"
        st.rerun()

def parse_slide_content(input_string):
    slides = re.split(r'Slide \d+:', input_string)[1:]
    parsed_slides = []
    
    for slide in slides:
        title_match = re.search(r'Title:(.*?)Body:', slide, re.DOTALL)
        body_match = re.search(r'Body:(.*?)(?=Slide \d+:|$)', slide, re.DOTALL)
        
        if title_match and body_match:
            title = title_match.group(1).strip()
            body = body_match.group(1).strip()
            parsed_slides.append((title, body))
    
    return parsed_slides

def create_slide(prs, title, content, background_image_path):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    # Set background image
    left = top = Inches(0)
    pic = slide.shapes.add_picture(background_image_path, left, top, width=prs.slide_width, height=prs.slide_height)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)
    
    # Set title
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(40)
    
    # Set content
    content_shape = slide.placeholders[1]
    content_shape.text = content
    for paragraph in content_shape.text_frame.paragraphs:
        paragraph.font.size = Pt(24)
    
    return slide

def create_presentation(input_string, background_image_path):
    prs = Presentation()
    slides_content = parse_slide_content(input_string)
    
    for title, body in slides_content:
        slide = create_slide(prs, title, body, background_image_path)
    
    # Save the presentation to a BytesIO object
    ppt_io = io.BytesIO()
    prs.save(ppt_io)
    ppt_io.seek(0)
    
    return ppt_io
    